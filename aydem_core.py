"""
Aydem - Günlük Üretim ve Ciro Karşılaştırması (2025 vs 2026) + Dönem Ortalamaları
===================================================================================

EPİAŞ Şeffaflık Platformu'ndan (eptr2 kütüphanesi ile) Aydem'in tüm
santralleri için 1 Ekim 2025'ten DÜNE kadar olan SAATLİK üretimi çeker,
GÜNLÜK toplamlara çevirir ve şunları üretir:

  1) 2026 Ekim (1 Ekim -> dün) ile 2025'in aynı günlerinin günlük karşılaştırması
  2) Ortalama günlük üretim ve ciro tablosu:
        - 1 Ekim 2025 -> dün (tüm dönem)
        - Ekim 2025, Kasım 2025, Aralık 2025
        - 2026 Ekim (1 Ekim -> dün)
     ve 2026 Ekim'in diğer dönemlere ORANI

Ciro hesabı (santral bazında):
  - YEKDEM kapsamındaki santral (YEKDEM_PLANTS):
        üretim (MWh) x YEKDEM fiyatı (USc/kWh) x 10 x USD/TRY kuru
  - Diğer santraller: saatlik üretim (MWh) x saatlik PTF (TL/MWh)

Çıktılar:
  - aydem_gunluk_karsilastirma.xlsx
        * Gunluk_Karsilastirma : günlük toplam üretim/ciro, iki yıl yan yana,
                                 % değişim (+ Excel içinde grafikler)
        * Ortalamalar          : dönem ortalamaları ve 2026 Ekim oranları
        * Ozet                 : günlük karşılaştırma dönem toplamları
        * Santral_Gunluk       : santral x gün detay (tüm dönem)
  - aydem_gunluk_karsilastirma.png : grafikler + ortalama tablosu

KURULUM
-------
    py -m pip install "eptr2[allextras]" matplotlib openpyxl

ÇALIŞTIRMA (cmd)
----------------
    set EPTR_USERNAME=mailadresin@ornek.com
    set EPTR_PASSWORD=sifren
    py aydem_gunluk_karsilastirma.py

NOT: Veri ~13 ay boyunca aylık parçalar halinde çekildiği için (18 santral x
~13 ay) çalışması birkaç on dakika sürebilir. Pencereyi kapatma.
"""

from __future__ import annotations

import calendar
import os
import sys
import unicodedata
from datetime import date, timedelta

import pandas as pd
from eptr2 import EPTR2

# ----------------------------------------------------------------------
# CONFIG
# ----------------------------------------------------------------------

EPIAS_USERNAME = os.environ.get("EPTR_USERNAME", "")
EPIAS_PASSWORD = os.environ.get("EPTR_PASSWORD", "")

AYDEM_PLANT_NAMES = [
    "AKINCI HES", "UŞAK RES", "GÖKTAŞ 2 HES", "KOYULHİSAR HES", "GÖKTAŞ 1 HES",
    "SÖKE RES", "YALOVA RES", "ÇIRAKDAMI HES", "TOROS HES", "DERELİ HES",
    "AKSU HES", "DALAMAN HES", "MENTAŞ HES", "KEMER HES", "GÖKYAR HES",
    "ADIGÜZEL HES", "BEREKET 1-2 HES", "FESLEK HES",
]

# YEKDEM kapsamındaki santraller ve birim fiyatları (USc/kWh).
# Listede OLMAYAN santraller PTF ile hesaplanır.
YEKDEM_PLANTS = {
    "AKINCI HES": 6.30,
    "SÖKE RES": 4.95,
    "YALOVA RES": 4.95,
}

# Karşılaştırma dönemi: her yıl için 1 Ekim -> dünün ay/günü.
START_MONTH, START_DAY = 10, 1
END_DATE_OVERRIDE = None   # None ise bitiş = DÜN. Sabitlemek için örn. date(2026, 10, 17) yaz

# YEKDEM'deki santraller için USD/TRY kuru - KENDİN GÜNCELLE.
# Önce ay bazlı değere ("YYYY-AA"), yoksa yıl bazlı değere bakılır.
# Kur yıl içinde çok değişiyorsa ay bazlı girmen daha doğru sonuç verir.
USDTRY_RATE_BY_YEAR = {
    2026: 43.0,   # KENDİN GÜNCELLE
    2025: 41.0,   # KENDİN GÜNCELLE
}
USDTRY_RATE_BY_MONTH = {
    # "2025-11": 42.0,
    # "2025-12": 42.8,
}

OUTPUT_XLSX = "aydem_gunluk_karsilastirma.xlsx"
OUTPUT_PNG = "aydem_gunluk_karsilastirma.png"

TR_MONTHS = {
    1: "Ocak", 2: "Şubat", 3: "Mart", 4: "Nisan", 5: "Mayıs", 6: "Haziran",
    7: "Temmuz", 8: "Ağustos", 9: "Eylül", 10: "Ekim", 11: "Kasım", 12: "Aralık",
}


# ----------------------------------------------------------------------
# Yardımcılar
# ----------------------------------------------------------------------

def _normalize_tr(text: str) -> str:
    if text is None:
        return ""
    text = text.replace("İ", "I").replace("ı", "i")
    text = "".join(
        c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn"
    )
    return text.upper()


def _to_local_naive(series: pd.Series) -> pd.Series:
    """Tarih sütununu Türkiye saatine çevirip timezone bilgisini kaldırır."""
    s = pd.to_datetime(series)
    if getattr(s.dt, "tz", None) is not None:
        s = s.dt.tz_convert("Europe/Istanbul").dt.tz_localize(None)
    return s


def connect() -> EPTR2:
    if not EPIAS_USERNAME or not EPIAS_PASSWORD:
        sys.exit(
            "HATA: EPTR_USERNAME / EPTR_PASSWORD ortam değişkenleri boş. "
            "EPİAŞ Şeffaflık Platformu kullanıcı adı ve şifreni ayarla."
        )
    return EPTR2(username=EPIAS_USERNAME, password=EPIAS_PASSWORD)


RETRY_KW = dict(
    request_kwargs={"timeout": 30},
    retry_attempts=5,
    retry_backoff=2,
    retry_backoff_max=30,
    retry_jitter=0.5,
)


def find_aydem_plants(eptr) -> pd.DataFrame:
    df = pd.DataFrame(eptr.call("pp-list"))
    if df.empty:
        sys.exit("HATA: pp-list servisinden santral listesi boş döndü.")

    search_cols = [c for c in ("name", "shortName") if c in df.columns]
    if not search_cols:
        print("Mevcut sütunlar:", list(df.columns))
        sys.exit("HATA: 'name'/'shortName' sütunu bulunamadı.")

    targets = [_normalize_tr(n) for n in AYDEM_PLANT_NAMES]

    def _matches(row) -> bool:
        for col in search_cols:
            val = _normalize_tr(str(row[col]))
            if any(t in val or val in t for t in targets):
                return True
        return False

    out = df[df.apply(_matches, axis=1)].copy()
    if out.empty:
        print(df[search_cols].head(30).to_string(index=False))
        sys.exit("HATA: AYDEM_PLANT_NAMES ile eşleşen santral bulunamadı.")

    print(f"\nBulunan Aydem santralleri ({len(out)} adet):")
    print(out[search_cols + ["id"]].to_string(index=False))
    return out


def get_pp_id_column(df: pd.DataFrame) -> str:
    for c in df.columns:
        if c.lower() in ("id", "powerplantid", "pp_id"):
            return c
    sys.exit(f"HATA: Santral ID sütunu bulunamadı. Sütunlar: {list(df.columns)}")


def yekdem_price_for(plant_name: str):
    norm = _normalize_tr(plant_name)
    for yek_name, price in YEKDEM_PLANTS.items():
        n = _normalize_tr(yek_name)
        if n in norm or norm in n:
            return price
    return None


def usdtry_for(d: date):
    key = f"{d.year}-{d.month:02d}"
    if key in USDTRY_RATE_BY_MONTH:
        return USDTRY_RATE_BY_MONTH[key]
    return USDTRY_RATE_BY_YEAR.get(d.year)


def fetch_hourly_production(eptr, pp_id, start_date: str, end_date: str) -> pd.DataFrame:
    df = pd.DataFrame(
        eptr.call("rt-gen", start_date=start_date, end_date=end_date, pp_id=pp_id, **RETRY_KW)
    )
    if df.empty:
        return df

    date_cols = [c for c in df.columns if "date" in c.lower()]
    if not date_cols:
        sys.exit(f"HATA: Tarih sütunu bulunamadı. Sütunlar: {list(df.columns)}")
    date_col = date_cols[0]
    df[date_col] = _to_local_naive(df[date_col])

    if "total" not in df.columns:
        num_cols = [c for c in df.columns if c != date_col and pd.api.types.is_numeric_dtype(df[c])]
        df["total"] = df[num_cols].sum(axis=1)

    return df[[date_col, "total"]].rename(columns={date_col: "tarih_saat", "total": "uretim_mwh"})


def fetch_ptf(eptr, start_date: str, end_date: str) -> pd.DataFrame:
    df = pd.DataFrame(eptr.call("mcp", start_date=start_date, end_date=end_date, **RETRY_KW))
    date_cols = [c for c in df.columns if "date" in c.lower()]
    price_cols = [
        c for c in df.columns
        if "price" in c.lower() or "ptf" in c.lower() or "mcp" in c.lower()
    ]
    if not date_cols or not price_cols:
        sys.exit(f"HATA: mcp servisi beklenmedik sütunlarla döndü: {list(df.columns)}")

    df["tarih_saat"] = _to_local_naive(df[date_cols[0]])
    df = df.rename(columns={price_cols[0]: "ptf_tl_mwh"})
    return df[["tarih_saat", "ptf_tl_mwh"]]


def month_chunks(start: date, end: date):
    """[start, end] aralığını ay sınırlarında parçalara böler."""
    chunks, cur = [], start
    while cur <= end:
        nxt = date(cur.year + (cur.month == 12), cur.month % 12 + 1, 1)
        chunks.append((cur, min(nxt - timedelta(days=1), end)))
        cur = nxt
    return chunks


# ----------------------------------------------------------------------
# Veri toplama
# ----------------------------------------------------------------------

def collect_daily(eptr, plants: pd.DataFrame, start: date, end: date):
    """Tüm santraller için start..end arası günlük üretim ve ciro (aylık parçalarla çeker)."""
    id_col = get_pp_id_column(plants)
    name_col = "name" if "name" in plants.columns else plants.columns[0]
    chunks = month_chunks(start, end)

    # Kur eksikse çekime başlamadan dur (YEKDEM'li santral varsa kur şart)
    for c_start, _ in chunks:
        if usdtry_for(c_start) is None:
            sys.exit(
                f"HATA: {c_start:%Y-%m} için USD/TRY kuru yok. "
                "USDTRY_RATE_BY_YEAR veya USDTRY_RATE_BY_MONTH içine ekle."
            )

    rows, failures = [], []
    for ci, (c_start, c_end) in enumerate(chunks, 1):
        s, e = c_start.isoformat(), c_end.isoformat()
        usdtry = usdtry_for(c_start)
        print(f"\n=== [{ci}/{len(chunks)}] {s} -> {e} ===")
        try:
            ptf = fetch_ptf(eptr, s, e)
        except Exception as exc:  # noqa: BLE001
            print(f"  HATA: PTF çekilemedi ({exc}). PTF'li santraller bu parça için atlanacak.")
            failures.append(f"PTF {s}->{e}")
            ptf = None

        for _, p in plants.iterrows():
            pname, pp_id = p[name_col], p[id_col]
            yek_price = yekdem_price_for(pname)
            mekanizma = "YEKDEM" if yek_price is not None else "PTF"
            print(f"  -> {pname} ({mekanizma})")

            if mekanizma == "PTF" and ptf is None:
                failures.append(f"{pname} {s}->{e} (PTF yok)")
                continue
            try:
                prod = fetch_hourly_production(eptr, pp_id, s, e)
            except Exception as exc:  # noqa: BLE001
                print(f"     HATA: veri çekilemedi ({exc}). Atlanıyor.")
                failures.append(f"{pname} {s}->{e}")
                continue
            if prod.empty:
                print("     UYARI: veri yok, atlanıyor.")
                failures.append(f"{pname} {s}->{e} (boş)")
                continue

            if mekanizma == "YEKDEM":
                prod["ciro_tl"] = prod["uretim_mwh"] * (yek_price / 100) * 1000 * usdtry
            else:
                prod = pd.merge(prod, ptf, on="tarih_saat", how="left")
                prod["ciro_tl"] = prod["uretim_mwh"] * prod["ptf_tl_mwh"].fillna(0)

            prod["tarih"] = prod["tarih_saat"].dt.normalize()
            daily = prod.groupby("tarih", as_index=False)[["uretim_mwh", "ciro_tl"]].sum()
            daily["santral"] = pname
            daily["mekanizma"] = mekanizma
            rows.append(daily)

    if not rows:
        sys.exit("HATA: Hiçbir santral/dönem için veri üretilemedi.")

    # Aynı gün birden fazla parçada gelirse (sınır saatleri) tek satıra indir
    detail = pd.concat(rows, ignore_index=True)
    detail = detail.groupby(["tarih", "santral", "mekanizma"], as_index=False)[
        ["uretim_mwh", "ciro_tl"]
    ].sum()
    return detail, failures


def comparison_window(detail: pd.DataFrame, end_day: date) -> pd.DataFrame:
    """Bu yıl ve geçen yıl için 1 Ekim -> bitiş günü penceresini seçer (gün hizalı)."""
    parts = []
    for year in (end_day.year, end_day.year - 1):
        s = pd.Timestamp(year, START_MONTH, START_DAY)
        last_dom = calendar.monthrange(year, end_day.month)[1]
        e = pd.Timestamp(year, end_day.month, min(end_day.day, last_dom))
        part = detail[(detail["tarih"] >= s) & (detail["tarih"] <= e)].copy()
        part["yil"] = year
        part["gun"] = (part["tarih"] - s).dt.days + 1
        parts.append(part)
    return pd.concat(parts, ignore_index=True)


# ----------------------------------------------------------------------
# Hesaplamalar
# ----------------------------------------------------------------------

def build_comparison(detail: pd.DataFrame, this_year: int):
    last_year = this_year - 1
    tot = detail.groupby(["gun", "yil"])[["uretim_mwh", "ciro_tl"]].sum().unstack("yil")

    comp = pd.DataFrame(index=tot.index)
    comp.index.name = "gun"
    for y in (last_year, this_year):
        comp[f"uretim_mwh_{y}"] = tot[("uretim_mwh", y)] if ("uretim_mwh", y) in tot else float("nan")
    comp["uretim_degisim"] = (
        comp[f"uretim_mwh_{this_year}"] / comp[f"uretim_mwh_{last_year}"] - 1
    ).replace([float("inf"), float("-inf")], float("nan"))
    for y in (last_year, this_year):
        comp[f"ciro_tl_{y}"] = tot[("ciro_tl", y)] if ("ciro_tl", y) in tot else float("nan")
    comp["ciro_degisim"] = (
        comp[f"ciro_tl_{this_year}"] / comp[f"ciro_tl_{last_year}"] - 1
    ).replace([float("inf"), float("-inf")], float("nan"))
    comp = comp.reset_index()

    comp.insert(1, "tarih_etiketi", comp["gun"].apply(
        lambda g: (date(this_year, START_MONTH, START_DAY) + timedelta(days=int(g) - 1)).strftime("%d.%m")
    ))
    return comp


def build_summary(comp: pd.DataFrame, this_year: int) -> pd.DataFrame:
    ly = this_year - 1
    u_this, u_last = comp[f"uretim_mwh_{this_year}"].sum(), comp[f"uretim_mwh_{ly}"].sum()
    c_this, c_last = comp[f"ciro_tl_{this_year}"].sum(), comp[f"ciro_tl_{ly}"].sum()

    def ratio(a, b):
        return a / b if b else float("nan")

    return pd.DataFrame(
        {
            "metrik": ["Toplam üretim (MWh)", "Toplam ciro (TL)"],
            str(ly): [u_last, c_last],
            str(this_year): [u_this, c_this],
            f"{this_year} / {ly} (%)": [ratio(u_this, u_last), ratio(c_this, c_last)],
            "değişim (%)": [ratio(u_this, u_last) - 1, ratio(c_this, c_last) - 1],
        }
    )


def build_averages(detail_all: pd.DataFrame, end_day: date) -> pd.DataFrame:
    """Dönem bazlı ORTALAMA GÜNLÜK üretim/ciro ve son dönemin (bu yıl Ekim) diğerlerine oranı."""
    ly, ty = end_day.year - 1, end_day.year
    daily = detail_all.groupby("tarih", as_index=False)[["uretim_mwh", "ciro_tl"]].sum()

    periods = [
        (f"{date(ly, START_MONTH, START_DAY):%d.%m.%Y} → {end_day:%d.%m.%Y}",
         date(ly, START_MONTH, START_DAY), end_day),
        (f"Ekim {ly}", date(ly, 10, 1), date(ly, 10, 31)),
        (f"Kasım {ly}", date(ly, 11, 1), date(ly, 11, 30)),
        (f"Aralık {ly}", date(ly, 12, 1), date(ly, 12, 31)),
        (f"{TR_MONTHS[START_MONTH]} {ty} (1 {TR_MONTHS[START_MONTH]} → {end_day:%d.%m})",
         date(ty, START_MONTH, START_DAY), end_day),
    ]

    rows = []
    for label, s, e in periods:
        e = min(e, end_day)
        sel = daily[(daily["tarih"] >= pd.Timestamp(s)) & (daily["tarih"] <= pd.Timestamp(e))]
        n = len(sel)
        rows.append(
            {
                "donem": label,
                "gun": n,
                "ort_uretim_mwh": sel["uretim_mwh"].mean() if n else float("nan"),
                "ort_ciro_tl": sel["ciro_tl"].mean() if n else float("nan"),
            }
        )
    avg = pd.DataFrame(rows)

    ref_u, ref_c = avg.iloc[-1]["ort_uretim_mwh"], avg.iloc[-1]["ort_ciro_tl"]
    inf = [float("inf"), float("-inf")]
    avg["oran_uretim"] = (ref_u / avg["ort_uretim_mwh"]).replace(inf, float("nan"))
    avg["oran_ciro"] = (ref_c / avg["ort_ciro_tl"]).replace(inf, float("nan"))
    avg.loc[avg.index[-1], ["oran_uretim", "oran_ciro"]] = float("nan")  # kendisiyle kıyas yok
    return avg


# ----------------------------------------------------------------------
# Çıktılar
# ----------------------------------------------------------------------

def write_excel(comp, detail_all, summary, avg, this_year):
    from openpyxl.chart import BarChart, LineChart, Reference

    ly = this_year - 1
    detail_out = detail_all.copy()
    detail_out["tarih"] = detail_out["tarih"].dt.date
    detail_out = detail_out[["tarih", "santral", "mekanizma", "uretim_mwh", "ciro_tl"]]
    detail_out = detail_out.sort_values(["tarih", "santral"])

    ref_label = f"{this_year} Ekim"
    avg_out = avg.rename(
        columns={
            "donem": "Dönem",
            "gun": "Gün sayısı",
            "ort_uretim_mwh": "Ort. günlük üretim (MWh)",
            "ort_ciro_tl": "Ort. günlük ciro (TL)",
            "oran_uretim": f"{ref_label} / dönem - üretim",
            "oran_ciro": f"{ref_label} / dönem - ciro",
        }
    )

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        comp.to_excel(writer, sheet_name="Gunluk_Karsilastirma", index=False)
        avg_out.to_excel(writer, sheet_name="Ortalamalar", index=False)
        summary.to_excel(writer, sheet_name="Ozet", index=False)
        detail_out.to_excel(writer, sheet_name="Santral_Gunluk", index=False)

        ws = writer.sheets["Gunluk_Karsilastirma"]
        n = len(comp)
        cols = {name: i + 1 for i, name in enumerate(comp.columns)}

        def add_chart(title, y_title, c1, c2, anchor):
            ch = LineChart()
            ch.title, ch.height, ch.width = title, 8, 18
            ch.y_axis.title, ch.x_axis.title = y_title, "Ekim (gün)"
            for col in (c1, c2):
                ch.add_data(Reference(ws, min_col=cols[col], min_row=1, max_row=n + 1), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=cols["tarih_etiketi"], min_row=2, max_row=n + 1))
            ws.add_chart(ch, anchor)

        add_chart("Günlük Üretim (MWh)", "MWh", f"uretim_mwh_{ly}", f"uretim_mwh_{this_year}", "K2")
        add_chart("Günlük Ciro (TL)", "TL", f"ciro_tl_{ly}", f"ciro_tl_{this_year}", "K20")

        bc = BarChart()
        bc.title, bc.height, bc.width = f"Günlük Değişim: {this_year} vs {ly}", 8, 18
        bc.y_axis.number_format = "0%"
        bc.x_axis.title = "Ekim (gün)"
        for col in ("uretim_degisim", "ciro_degisim"):
            bc.add_data(Reference(ws, min_col=cols[col], min_row=1, max_row=n + 1), titles_from_data=True)
        bc.set_categories(Reference(ws, min_col=cols["tarih_etiketi"], min_row=2, max_row=n + 1))
        ws.add_chart(bc, "K38")

        for name, idx in cols.items():
            fmt = "0.0%" if name.endswith("degisim") else ("#,##0" if name.startswith(("uretim", "ciro")) else None)
            if fmt:
                for r in range(2, n + 2):
                    ws.cell(row=r, column=idx).number_format = fmt

        ws_s = writer.sheets["Ozet"]
        for r in (2, 3):
            ws_s.cell(row=r, column=2).number_format = "#,##0"
            ws_s.cell(row=r, column=3).number_format = "#,##0"
            ws_s.cell(row=r, column=4).number_format = "0.0%"
            ws_s.cell(row=r, column=5).number_format = "0.0%"

        ws_a = writer.sheets["Ortalamalar"]
        for r in range(2, len(avg_out) + 2):
            ws_a.cell(row=r, column=3).number_format = "#,##0"
            ws_a.cell(row=r, column=4).number_format = "#,##0"
            ws_a.cell(row=r, column=5).number_format = "0.0%"
            ws_a.cell(row=r, column=6).number_format = "0.0%"

        for sheet in writer.sheets.values():
            for col_cells in sheet.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
                sheet.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 40)

    print(f"\nExcel kaydedildi: {OUTPUT_XLSX}")


def write_chart(comp, avg, this_year):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    ly = this_year - 1
    u_this, u_last = comp[f"uretim_mwh_{this_year}"].sum(), comp[f"uretim_mwh_{ly}"].sum()
    c_this, c_last = comp[f"ciro_tl_{this_year}"].sum(), comp[f"ciro_tl_{ly}"].sum()
    u_ratio = u_this / u_last if u_last else float("nan")
    c_ratio = c_this / c_last if c_last else float("nan")

    def pct_text(ratio):
        return f"{this_year}, {ly}'in %{ratio * 100:.1f}'i  ({(ratio - 1) * 100:+.1f}%)"

    labels = comp["tarih_etiketi"]
    fig = plt.figure(figsize=(12, 18))
    gs = fig.add_gridspec(4, 1, height_ratios=[1, 1, 1, 1.05], hspace=0.5)
    ax1 = fig.add_subplot(gs[0])
    ax2 = fig.add_subplot(gs[1], sharex=ax1)
    ax3 = fig.add_subplot(gs[2], sharex=ax1)
    ax4 = fig.add_subplot(gs[3])
    plt.setp(ax1.get_xticklabels(), visible=False)
    plt.setp(ax2.get_xticklabels(), visible=False)

    ax1.plot(labels, comp[f"uretim_mwh_{ly}"], marker="o", label=str(ly), color="#888888")
    ax1.plot(labels, comp[f"uretim_mwh_{this_year}"], marker="o", label=str(this_year), color="#2b7de9")
    ax1.set_ylabel("Üretim (MWh)")
    ax1.set_title(f"Günlük Toplam Üretim\nToplam üretim: {pct_text(u_ratio)}", fontsize=12)
    ax1.grid(alpha=0.3)
    ax1.legend()

    ax2.plot(labels, comp[f"ciro_tl_{ly}"], marker="o", label=str(ly), color="#888888")
    ax2.plot(labels, comp[f"ciro_tl_{this_year}"], marker="o", label=str(this_year), color="#2b7de9")
    ax2.set_ylabel("Ciro (milyon TL)")
    ax2.yaxis.set_major_formatter(lambda v, _: f"{v / 1e6:.1f}")
    ax2.set_title(f"Günlük Ciro (YEKDEM santraller YEKDEM, diğerleri PTF)\nToplam ciro: {pct_text(c_ratio)}", fontsize=12)
    ax2.grid(alpha=0.3)
    ax2.legend()

    x = np.arange(len(comp))
    w = 0.4
    ax3.bar(x - w / 2, comp["uretim_degisim"] * 100, w, label="Üretim değişimi", color="#43a047")
    ax3.bar(x + w / 2, comp["ciro_degisim"] * 100, w, label="Ciro değişimi", color="#fb8c00")
    ax3.axhline(0, color="black", linewidth=0.8)
    ax3.set_ylabel(f"{this_year} / {ly} değişim (%)")
    ax3.yaxis.set_major_formatter(lambda v, _: f"{v:+.0f}%")
    ax3.set_title(f"Günlük Değişim: {this_year}, {ly}'e göre %", fontsize=12)
    ax3.set_xticks(x)
    ax3.set_xticklabels(labels, rotation=45)
    ax3.grid(alpha=0.3, axis="y")
    ax3.legend()

    # ---------------- Ortalama tablosu ----------------
    ax4.axis("off")
    ax4.set_title(
        f"Ortalama Günlük Üretim ve Ciro  |  {this_year} Ekim'in önceki dönemlere oranı",
        fontsize=13, fontweight="bold", pad=36,
    )

    def ratio_text(r):
        if pd.isna(r):
            return "—"
        return f"%{r * 100:.1f}\n({(r - 1) * 100:+.1f}%)"

    col_labels = [
        "Dönem", "Gün\nsayısı", "Ort. günlük\nüretim (MWh)", "Ort. günlük\nciro (milyon TL)",
        f"{this_year} Ekim\n÷ dönem\nÜRETİM", f"{this_year} Ekim\n÷ dönem\nCİRO",
    ]
    cell_text = []
    for _, r in avg.iterrows():
        cell_text.append([
            r["donem"],
            f"{int(r['gun'])}",
            "—" if pd.isna(r["ort_uretim_mwh"]) else f"{r['ort_uretim_mwh']:,.0f}",
            "—" if pd.isna(r["ort_ciro_tl"]) else f"{r['ort_ciro_tl'] / 1e6:,.2f}",
            ratio_text(r["oran_uretim"]),
            ratio_text(r["oran_ciro"]),
        ])

    table = ax4.table(
        cellText=cell_text, colLabels=col_labels, cellLoc="center", loc="upper center",
        colWidths=[0.30, 0.07, 0.16, 0.17, 0.15, 0.15],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(10.5)
    table.scale(1, 2.9)

    last_row = len(cell_text)
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor("#c9ced6")
        if row == 0:
            cell.set_facecolor("#1f3a5f")
            cell.set_text_props(color="white", fontweight="bold")
            continue
        base = "#f4f6f9" if row % 2 == 0 else "white"
        cell.set_facecolor(base)
        if col == 0:
            cell.set_text_props(ha="left")
            cell._loc = "left"
        if row == last_row:
            cell.set_facecolor("#dbe9ff")
            cell.set_text_props(fontweight="bold")
        if col in (4, 5) and row != last_row:
            val = avg.iloc[row - 1]["oran_uretim" if col == 4 else "oran_ciro"]
            if not pd.isna(val):
                cell.set_facecolor("#d9f2dd" if val >= 1 else "#fbdcdc")
                cell.set_text_props(color="#1b5e20" if val >= 1 else "#b71c1c", fontweight="bold")

    ax4.text(
        0.5, 1.01,
        "Ortalamalar günlük toplamların ortalamasıdır. Yeşil: Ekim ortalaması o dönemden yüksek, kırmızı: düşük.",
        ha="center", va="bottom", fontsize=9, color="#555555", transform=ax4.transAxes,
    )

    fig.savefig(OUTPUT_PNG, dpi=150, bbox_inches="tight")
    plt.close(fig)
    print(f"Grafik kaydedildi: {OUTPUT_PNG}")


def main():
    end_day = END_DATE_OVERRIDE or (date.today() - timedelta(days=1))  # bugünün verisi yok -> dün
    this_year = end_day.year
    if end_day < date(this_year, START_MONTH, START_DAY):
        sys.exit("HATA: Bitiş tarihi 1 Ekim'den önce; bu program Ekim ayı karşılaştırması için yazıldı.")
    start_day = date(this_year - 1, START_MONTH, START_DAY)
    print(f"Veri dönemi: {start_day} -> {end_day}")
    print(f"Günlük karşılaştırma: 1 Ekim -> {end_day:%d.%m} ({this_year - 1} vs {this_year})")

    eptr = connect()
    print("\n1) Aydem santralleri aranıyor (pp-list)...")
    plants = find_aydem_plants(eptr)

    print("\n2) Veriler çekiliyor (aylık parçalar)...")
    detail_all, failures = collect_daily(eptr, plants, start_day, end_day)

    comp = build_comparison(comparison_window(detail_all, end_day), this_year)
    summary = build_summary(comp, this_year)
    avg = build_averages(detail_all, end_day)

    write_excel(comp, detail_all, summary, avg, this_year)
    write_chart(comp, avg, this_year)

    print("\n===== GÜNLÜK KARŞILAŞTIRMA =====")
    print(comp.to_string(index=False, float_format=lambda v: f"{v:,.3f}"))

    print("\n===== ÖZET (günlük karşılaştırma toplamları) =====")
    ly = this_year - 1
    for _, r in summary.iterrows():
        print(
            f"{r['metrik']}: {ly} = {r[str(ly)]:,.0f} | {this_year} = {r[str(this_year)]:,.0f} "
            f"-> {this_year}, {ly}'in %{r.iloc[3] * 100:.1f}'i (değişim {r.iloc[4] * 100:+.1f}%)"
        )

    print("\n===== ORTALAMA GÜNLÜK ÜRETİM VE CİRO =====")
    for _, r in avg.iterrows():
        line = (
            f"{r['donem']:<34} {int(r['gun']):>4} gün | üretim {r['ort_uretim_mwh']:>10,.0f} MWh"
            f" | ciro {r['ort_ciro_tl'] / 1e6:>8,.2f} M TL"
        )
        if not pd.isna(r["oran_uretim"]):
            line += f" | {this_year} Ekim oranı: üretim %{r['oran_uretim'] * 100:.1f}, ciro %{r['oran_ciro'] * 100:.1f}"
        print(line)

    if failures:
        print(f"\nUYARI: {len(failures)} parça çekilemedi/boş geldi, ortalamalar bu yüzden etkilenmiş olabilir:")
        for f in failures[:30]:
            print("  -", f)


if __name__ == "__main__":
    main()
